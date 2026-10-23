import os, sys, csv
sys.path.insert(0, os.path.join(os.path.dirname(__file__), ".."))
import pytest
from desktop.database.db_manager import DBManager, DuplicateError
from desktop.services import export_service as ex
from desktop.services import import_backup_service as imp

@pytest.fixture()
def db(tmp_path):
    d = DBManager(str(tmp_path / "t.db")); yield d; d.close()

def test_profile_crud_and_duplicate(db):
    pid = db.add_profile("Ali Khan", facebook_user_id="10001", profile_url="https://facebook.com/ali")
    assert pid
    with pytest.raises(DuplicateError): db.add_profile("Copy", facebook_user_id="10001")
    db.update_profile(pid, notes="vip", connection_status="Connected")
    rows = db.list_profiles(search="Ali"); assert len(rows) == 1 and rows[0]["notes"] == "vip"
    assert db.list_profiles(search="zzz") == []
    db.delete_profile(pid); assert db.list_profiles() == []

def test_page_crud_duplicate_and_link(db):
    pid = db.add_profile("Owner One")
    pg = db.add_page("555", "My Page", page_url="https://facebook.com/mypage", category="Business")
    with pytest.raises(DuplicateError): db.add_page("555", "Dup")
    db.link_profile_page(pid, pg, access_role="Admin")
    rows = db.list_pages(search="My"); assert rows[0]["profiles"] == "Owner One"
    assert db.profile_page_count(pid) == 1
    assert db.list_pages(profile_id=pid)[0]["facebook_page_id"] == "555"
    assert db.list_pages(manual_status="Active") == []

def test_missing_pages_flagged_not_deleted(db):
    pid = db.add_profile("P"); pg1 = db.add_page("1", "One"); pg2 = db.add_page("2", "Two")
    db.link_profile_page(pid, pg1); db.link_profile_page(pid, pg2)
    missing = db.mark_missing_pages(pid, {"1"})
    assert missing == [pg2]
    rows = {r["facebook_page_id"]: r for r in db.list_pages()}
    assert rows["2"]["manual_status"] == "Missing - Verify" and len(rows) == 2

def test_health_and_monetization_unknown_defaults(db):
    pg = db.add_page("9", "Health Page")
    h = db.latest_health(pg); assert h["technical_connection_status"] == "Unknown"
    m = db.latest_monetization(pg); assert m["api_status"] == "Unknown" and m["eligibility_status"] == "Unknown"
    db.record_health(pg, technical="Healthy", policy="Unknown", restriction="Unknown", source="api")
    assert db.latest_health(pg)["technical_connection_status"] == "Healthy"
    db.record_monetization(pg, api_status="Unknown", manual_status="Manually Verified")
    assert db.latest_monetization(pg)["manual_status"] == "Manually Verified"

def test_dashboard_stats_real_data(db):
    db.add_profile("A", connection_status="Connected"); db.add_profile("B")
    db.add_page("1", "One"); db.add_page("2", "Two")
    s = db.dashboard_stats()
    assert s["total_profiles"] == 2 and s["connected_profiles"] == 1 and s["total_pages"] == 2
    assert s["monetization_unknown"] == 2

def test_users_and_roles(db):
    db.create_user("aamir", "StrongPass!1", role="Super Admin")
    assert db.has_super_admin()
    assert db.verify_user("aamir", "StrongPass!1")["role"] == "Super Admin"
    assert db.verify_user("aamir", "wrong") is None
    with pytest.raises(DuplicateError): db.create_user("aamir", "x")

def test_exports_csv_excel_pdf(db, tmp_path):
    db.add_page("7", "Export Page", page_url="https://facebook.com/x", follower_count=1200)
    cols, rows = ex.pages_dataset(db)
    csvp = str(tmp_path / "r.csv"); ex.export_csv(rows, cols, csvp)
    with open(csvp, encoding="utf-8-sig") as f: assert "Export Page" in f.read()
    with pytest.raises(FileExistsError): ex.export_csv(rows, cols, csvp)  # no silent overwrite
    xp = str(tmp_path / "r.xlsx"); ex.export_excel({"Pages": (cols, rows)}, xp)
    from openpyxl import load_workbook; ws = load_workbook(xp)["Pages"]
    assert ws.freeze_panes == "A2" and ws["B2"].value == "Export Page"
    pp = str(tmp_path / "r.pdf"); ex.export_pdf("Pages Report", cols, rows, pp, summary="Total: 1")
    assert os.path.getsize(pp) > 500

def test_import_preview_and_duplicates(db, tmp_path):
    p = tmp_path / "imp.csv"
    p.write_text("facebook_page_id,page_name\n1,One\n1,DupInFile\n,NoId\n2,Two\n", encoding="utf-8")
    prev = imp.preview_csv(str(p))
    assert len(prev["valid"]) == 2 and len(prev["errors"]) == 2
    r = imp.import_pages(db, prev); assert r["imported"] == 2
    r2 = imp.import_pages(db, prev); assert r2["imported"] == 0 and len(r2["skipped"]) == 2

def test_backup_restore_requires_confirmation(db, tmp_path):
    db.add_page("3", "Backup Page"); src = db.db_path
    dest = str(tmp_path / "bk.db"); imp.backup_db(src, dest); assert os.path.exists(dest)
    with pytest.raises(PermissionError): imp.restore_db(src, dest, confirmed=False)
    db.add_page("4", "After Backup")
    imp.restore_db(src, dest, confirmed=True)
    db2 = DBManager(src); assert len(db2.list_pages()) == 1; db2.close()
