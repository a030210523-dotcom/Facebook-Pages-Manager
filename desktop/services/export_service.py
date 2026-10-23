"""CSV / Excel / PDF exports (Section 12). Never exports tokens, passwords or secrets."""
import csv, os
from datetime import datetime

PAGE_COLUMNS = ["facebook_page_id","page_name","page_url","category","follower_count",
                "manual_status","data_source","last_checked_at","notes","profiles"]
PROFILE_COLUMNS = ["facebook_user_id","display_name","profile_url","email","phone",
                   "connection_status","authorization_status","last_sync_at","notes"]

def _check_overwrite(path):
    if os.path.exists(path):
        raise FileExistsError(f"File already exists, choose another name: {path}")

def export_csv(rows, columns, path):
    _check_overwrite(path)
    with open(path, "w", newline="", encoding="utf-8-sig") as f:
        w = csv.DictWriter(f, fieldnames=columns, extrasaction="ignore")
        w.writeheader()
        for r in rows: w.writerow({c: r.get(c) for c in columns})
    return path

def export_excel(datasets, path):
    """datasets: {sheet_name: (columns, rows)}"""
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter
    _check_overwrite(path)
    wb = Workbook(); wb.remove(wb.active)
    for sheet, (columns, rows) in datasets.items():
        ws = wb.create_sheet(sheet[:31])
        ws.append(columns)
        for r in rows: ws.append([r.get(c) for c in columns])
        ws.freeze_panes = "A2"; ws.auto_filter.ref = ws.dimensions
        for i, c in enumerate(columns, 1):
            width = max([len(str(c))] + [len(str(r.get(c) or "")) for r in rows] + [10])
            ws.column_dimensions[get_column_letter(i)].width = min(width + 2, 60)
    wb.save(path); return path

def export_pdf(title, columns, rows, path, summary=None):
    from reportlab.lib.pagesizes import A4, landscape
    from reportlab.lib import colors
    from reportlab.platypus import SimpleDocTemplate, Table, TableStyle, Paragraph, Spacer
    from reportlab.lib.styles import getSampleStyleSheet
    _check_overwrite(path)
    styles = getSampleStyleSheet()
    def footer(canvas, doc):
        canvas.saveState(); canvas.setFont("Helvetica", 8)
        canvas.drawString(40, 20, "Facebook Pages Manager — Owner: Muhammad Aamir")
        canvas.drawRightString(A4[1]-40 if False else 800, 20, f"Page {doc.page}")
        canvas.restoreState()
    doc = SimpleDocTemplate(path, pagesize=landscape(A4), topMargin=36, bottomMargin=30)
    story = [Paragraph(title, styles["Title"]),
             Paragraph(f"Generated: {datetime.now().strftime('%Y-%m-%d %H:%M')}", styles["Normal"])]
    if summary:
        story.append(Paragraph(summary, styles["Normal"]))
    story.append(Spacer(1, 12))
    data = [columns] + [[str(r.get(c) if r.get(c) is not None else "")[:60] for c in columns] for r in rows]
    t = Table(data, repeatRows=1)
    t.setStyle(TableStyle([("BACKGROUND",(0,0),(-1,0),colors.HexColor("#1f4e79")),
                           ("TEXTCOLOR",(0,0),(-1,0),colors.white),("FONTSIZE",(0,0),(-1,-1),7),
                           ("GRID",(0,0),(-1,-1),0.4,colors.grey),("ROWBACKGROUNDS",(0,1),(-1,-1),[colors.white,colors.HexColor("#eef3f8")])]))
    story.append(t); doc.build(story, onFirstPage=footer, onLaterPages=footer); return path

def pages_dataset(db, **filters):
    rows = db.list_pages(**filters)
    for r in rows:
        h = db.latest_health(r["id"]); m = db.latest_monetization(r["id"])
        r["health_status"] = h["technical_connection_status"]; r["monetization_status"] = m["api_status"]
    cols = PAGE_COLUMNS + ["health_status","monetization_status"]
    return cols, rows
