"""Import preview/validation + DB backup/restore (Section 13)."""
import csv, os, shutil
from ..database.db_manager import DuplicateError

def preview_csv(path):
    with open(path, newline="", encoding="utf-8-sig") as f:
        rows = list(csv.DictReader(f))
    return _validate(rows)

def preview_excel(path, sheet=None):
    from openpyxl import load_workbook
    wb = load_workbook(path, read_only=True); ws = wb[sheet] if sheet else wb.active
    it = ws.iter_rows(values_only=True); headers = [str(h) for h in next(it)]
    rows = [dict(zip(headers, r)) for r in it]
    return _validate(rows)

def _validate(rows):
    valid, errors = [], []
    seen = set()
    for i, r in enumerate(rows, 2):
        pid = str(r.get("facebook_page_id") or "").strip()
        name = str(r.get("page_name") or "").strip()
        if not pid: errors.append(f"Row {i}: missing facebook_page_id"); continue
        if not name: errors.append(f"Row {i}: missing page_name"); continue
        if pid in seen: errors.append(f"Row {i}: duplicate facebook_page_id in file ({pid})"); continue
        seen.add(pid); valid.append(r)
    return {"valid": valid, "errors": errors, "total": len(rows)}

def import_pages(db, preview):
    imported, skipped = 0, []
    for r in preview["valid"]:
        try:
            db.add_page(str(r["facebook_page_id"]).strip(), str(r["page_name"]).strip(),
                        page_url=r.get("page_url"), category=r.get("category"),
                        follower_count=r.get("follower_count"), notes=r.get("notes"),
                        manual_status=r.get("manual_status") or "Unknown", data_source="import")
            imported += 1
        except DuplicateError as e: skipped.append(str(e))
    return {"imported": imported, "skipped": skipped}

def backup_db(db_path, dest_path):
    if os.path.exists(dest_path): raise FileExistsError(dest_path)
    shutil.copy2(db_path, dest_path); return dest_path

def restore_db(db_path, backup_path, confirmed=False):
    if not confirmed: raise PermissionError("Restore requires explicit confirmation")
    if not os.path.exists(backup_path): raise FileNotFoundError(backup_path)
    shutil.copy2(backup_path, db_path); return db_path
